#!/usr/bin/env python3
"""Build the deterministic, read-only F3 SacTown demo manifest."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl


DEFAULT_INPUT = Path("/Users/jkazmierski/Downloads/260917_F3 SacTown Data.xlsx")
HISTORICAL_START = date(2026, 8, 31)
HISTORICAL_END = date(2026, 9, 12)
OPTIONAL_PARTIAL_START = date(2026, 9, 14)
OPTIONAL_PARTIAL_END = date(2026, 9, 15)
SLOT_START = date(2026, 9, 28)
SLOT_END = date(2026, 10, 10)
REGION_KEY = "sactown-demo:region:v1"


def deterministic_uuid(source_key: str) -> str:
    digest = bytearray(hashlib.sha1(source_key.encode("utf-8")).digest()[:16])
    digest[6] = (digest[6] & 0x0F) | 0x50
    digest[8] = (digest[8] & 0x3F) | 0x80
    return str(uuid.UUID(bytes=bytes(digest)))


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize(value) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).lower())


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def parse_time_range(value: str) -> tuple[str, int]:
    match = re.fullmatch(
        r"\s*(\d{1,2}):(\d{2})(am|pm)\s*-\s*(\d{1,2}):(\d{2})(am|pm)\s*",
        clean(value).lower(),
    )
    if not match:
        raise ValueError(f"Unsupported source time range: {value!r}")

    def minutes(hour: str, minute: str, meridiem: str) -> int:
        h = int(hour) % 12
        if meridiem == "pm":
            h += 12
        return h * 60 + int(minute)

    start = minutes(match[1], match[2], match[3])
    end = minutes(match[4], match[5], match[6])
    if end <= start:
        end += 24 * 60
    return f"{start // 60:02d}:{start % 60:02d}", end - start


WEEKDAY = {
    "monday": 1,
    "tuesday": 2,
    "wednesday": 3,
    "thursday": 4,
    "friday": 5,
    "saturday": 6,
    "sunday": 0,
}


def load_rows(workbook, sheet_name: str, width: int):
    iterator = workbook[sheet_name].iter_rows(values_only=True)
    headers = next(iterator)
    rows = []
    for excel_row, raw in enumerate(iterator, start=2):
        row = tuple(raw) + (None,) * max(0, width - len(raw))
        row = row[:width]
        if any(value not in (None, "") for value in row):
            rows.append((excel_row, row))
    return headers, rows


def build_manifest(input_path: Path) -> dict:
    workbook = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    expected = {"Instructions", "Master Data", "Roster", "Sites", "Tower Data", "Lookup"}
    if set(workbook.sheetnames) != expected:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")

    _, master_rows = load_rows(workbook, "Master Data", 12)
    _, roster_rows = load_rows(workbook, "Roster", 6)
    _, site_rows = load_rows(workbook, "Sites", 26)

    roster_by_name = {}
    for excel_row, row in roster_rows:
        if row[0]:
            roster_by_name.setdefault(clean(row[0]), []).append((excel_row, row))

    selected_rows = []
    optional_rows = []
    for excel_row, row in master_rows:
        row_date = as_date(row[0])
        if not row_date or not row[1] or not row[2]:
            continue
        if HISTORICAL_START <= row_date <= HISTORICAL_END:
            if row[4] not in (None, 0, "") or row[6] not in (None, 0, ""):
                selected_rows.append((excel_row, row))
        elif OPTIONAL_PARTIAL_START <= row_date <= OPTIONAL_PARTIAL_END:
            if row[4] not in (None, 0, "") or row[6] not in (None, 0, ""):
                optional_rows.append((excel_row, row))

    names = sorted({clean(row[2]) for _, row in selected_rows}, key=str.casefold)
    normalized_groups = defaultdict(list)
    for name in names:
        normalized_groups[normalize(name)].append(name)
    collisions = [values for values in normalized_groups.values() if len(values) > 1]
    if collisions:
        raise ValueError(f"Selected-window normalized identity collisions: {collisions}")

    region_id = deterministic_uuid(REGION_KEY)
    project_id = deterministic_uuid("sactown-demo:import-project:v1")
    members = []
    member_by_name = {}
    for name in names:
        member_id = deterministic_uuid(f"sactown-demo:member:{normalize(name)}")
        appearances = [(excel_row, row) for excel_row, row in selected_rows if clean(row[2]) == name]
        source_regions = sorted({clean(row[2]) for _, row in roster_by_name.get(name, []) if row[2]})
        roster_matches = roster_by_name.get(name, [])
        first_seen = min(as_date(row[0]) for _, row in appearances)
        last_seen = max(as_date(row[0]) for _, row in appearances)
        member = {
            "id": member_id,
            "sourceIdentityId": deterministic_uuid(f"sactown-demo:source-identity:{normalize(name)}"),
            "sourceKey": f"member:{normalize(name)}",
            "paxName": name,
            "normalizedName": normalize(name),
            "firstSeenDate": first_seen.isoformat(),
            "lastSeenDate": last_seen.isoformat(),
            "sourceRegions": source_regions,
            "rosterRows": [excel_row for excel_row, _ in roster_matches],
        }
        members.append(member)
        member_by_name[name] = member

    active_site_rows = [(excel_row, row) for excel_row, row in site_rows if row[0] and not row[5]]
    active_ao_names = {clean(row[0]) for _, row in active_site_rows}
    historical_ao_names = {clean(row[1]) for _, row in selected_rows}
    if historical_ao_names != active_ao_names:
        raise ValueError(
            "Historical AO set does not exactly match active source schedules: "
            f"historical-only={sorted(historical_ao_names-active_ao_names)}, "
            f"schedule-only={sorted(active_ao_names-historical_ao_names)}"
        )

    sites_by_address = {}
    aos = []
    schedules = []
    for excel_row, row in active_site_rows:
        ao_name = clean(row[0])
        address = clean(row[15])
        location = clean(row[7])
        address_key = normalize(address)
        if not address_key:
            raise ValueError(f"Active AO {ao_name} has no source address")
        if address_key not in sites_by_address:
            site_source_key = f"site:{address_key}"
            sites_by_address[address_key] = {
                "id": deterministic_uuid(f"sactown-demo:{site_source_key}"),
                "sourceKey": site_source_key,
                "name": location,
                "address": address,
                "sourceRows": [],
                "recordClass": "imported_source_structure",
            }
        site = sites_by_address[address_key]
        site["sourceRows"].append(excel_row)

        start_time, duration = parse_time_range(row[8])
        weekday_name = clean(row[1]).lower()
        if weekday_name not in WEEKDAY:
            raise ValueError(f"Unsupported weekday for {ao_name}: {row[1]!r}")
        ao_source_key = f"ao:{normalize(ao_name)}"
        ao_id = deterministic_uuid(f"sactown-demo:{ao_source_key}")
        aos.append({
            "id": ao_id,
            "sourceKey": ao_source_key,
            "name": ao_name,
            "siteId": site["id"],
            "siteSourceKey": site["sourceKey"],
            "sourceRow": excel_row,
            "style": clean(row[9]),
            "recordClass": "imported_source_structure",
        })
        schedule_source_key = f"schedule:{normalize(ao_name)}:{WEEKDAY[weekday_name]}:{start_time}"
        schedules.append({
            "id": deterministic_uuid(f"sactown-demo:{schedule_source_key}"),
            "sourceKey": schedule_source_key,
            "aoId": ao_id,
            "aoSourceKey": ao_source_key,
            "siteId": site["id"],
            "siteSourceKey": site["sourceKey"],
            "weekday": WEEKDAY[weekday_name],
            "startTime": start_time,
            "durationMinutes": duration,
            "effectiveStartDate": as_date(row[13]).isoformat() if as_date(row[13]) else None,
            "sourceRow": excel_row,
            "recordClass": "imported_source_structure",
        })

    ao_by_name = {ao["name"]: ao for ao in aos}
    grouped = defaultdict(list)
    for excel_row, row in selected_rows:
        grouped[(as_date(row[0]).isoformat(), clean(row[1]))].append((excel_row, row))

    sessions = []
    raw_source_rows = []
    for (session_date, ao_name), rows in sorted(grouped.items()):
        ao = ao_by_name[ao_name]
        attendee_names = sorted({
            clean(row[2]) for _, row in rows
            if row[4] not in (None, 0, "") or row[6] not in (None, 0, "")
        }, key=str.casefold)
        post_names = sorted({clean(row[2]) for _, row in rows if row[4] not in (None, 0, "")}, key=str.casefold)
        qsource_names = sorted({clean(row[2]) for _, row in rows if row[6] not in (None, 0, "")}, key=str.casefold)
        q_names = sorted({
            clean(row[2]) for _, row in rows
            if row[5] not in (None, 0, "") or row[7] not in (None, 0, "")
        }, key=str.casefold)
        fng_names = sorted({clean(row[2]) for _, row in rows if row[3] not in (None, 0, "", "N", "n")}, key=str.casefold)
        schedule = next(s for s in schedules if s["aoId"] == ao["id"])
        source_key = f"session:{session_date}:{normalize(ao_name)}"
        session = {
            "id": deterministic_uuid(f"sactown-demo:{source_key}"),
            "stagedSessionId": deterministic_uuid(f"sactown-demo:staged:{source_key}"),
            "sourceKey": source_key,
            "date": session_date,
            "aoId": ao["id"],
            "aoSourceKey": ao["sourceKey"],
            "aoName": ao_name,
            "siteId": ao["siteId"],
            "siteSourceKey": ao["siteSourceKey"],
            "startTime": schedule["startTime"],
            "attendeeIds": [member_by_name[name]["id"] for name in attendee_names],
            "attendeeNames": attendee_names,
            "postNames": post_names,
            "qSourceNames": qsource_names,
            "qIds": [member_by_name[name]["id"] for name in q_names],
            "qNames": q_names,
            "fngs": [{
                "memberId": member_by_name[name]["id"],
                "paxName": name,
                "realName": "",
                "inviterIds": [],
                "invitedById": None,
            } for name in fng_names],
            "sourceRows": [excel_row for excel_row, _ in rows],
            "sourcePostFlags": len(post_names),
            "sourceQSourceFlags": len(qsource_names),
            "sourceAttendanceFlags": len(post_names) + len(qsource_names),
            "runtimeAttendanceRelationships": len(attendee_names),
            "recordClass": "imported_historical",
        }
        sessions.append(session)
        for excel_row, row in rows:
            raw_source_rows.append({
                "id": deterministic_uuid(f"sactown-demo:raw-row:{excel_row}"),
                "rowNumber": excel_row,
                "sourceKey": f"master-data-row:{excel_row}",
                "payload": {
                    "recordClass": "imported_historical",
                    "postDate": session_date,
                    "aoName": ao_name,
                    "paxName": clean(row[2]),
                    "fng": row[3],
                    "post": row[4],
                    "q": row[5],
                    "qSource": row[6],
                    "qQSource": row[7],
                },
            })

    prospective_slots = []
    current = SLOT_START
    while current <= SLOT_END:
        weekday = (current.weekday() + 1) % 7
        for schedule in schedules:
            if schedule["weekday"] != weekday:
                continue
            effective = schedule["effectiveStartDate"]
            if effective and current < date.fromisoformat(effective):
                continue
            source_key = (
                f"q-slot:{schedule['aoSourceKey']}:{current.isoformat()}:"
                f"{schedule['startTime']}"
            )
            prospective_slots.append({
                "id": deterministic_uuid(f"sactown-demo:{source_key}"),
                "sourceKey": source_key,
                "date": current.isoformat(),
                "aoId": schedule["aoId"],
                "siteId": schedule["siteId"],
                "startTime": schedule["startTime"],
                "durationMinutes": schedule["durationMinutes"],
                "qUserId": None,
                "scheduleSourceKey": schedule["sourceKey"],
                "recordClass": "generated_prospective",
            })
        current += timedelta(days=1)

    optional_groups = {(as_date(row[0]).isoformat(), clean(row[1])) for _, row in optional_rows}
    optional_names = {clean(row[2]) for _, row in optional_rows}
    manifest = {
        "manifestVersion": 1,
        "source": {
            "filename": input_path.name,
            "sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(),
            "sheet": "Master Data",
        },
        "region": {
            "id": region_id,
            "sourceKey": REGION_KEY,
            "name": "F3 SacTown Demo",
            "environment": "test",
            "lifecycleStatus": "active",
            "includeInReporting": False,
            "timezone": "America/Los_Angeles",
        },
        "importProject": {
            "id": project_id,
            "name": "F3 SacTown Demo — corrected historical source import",
            "sourceSystem": "260917_F3 SacTown Data.xlsx",
            "batchId": deterministic_uuid("sactown-demo:batch:master-data:v1"),
        },
        "historicalPolicy": {
            "startDate": HISTORICAL_START.isoformat(),
            "endDate": HISTORICAL_END.isoformat(),
            "datesTransposed": False,
            "recordClass": "imported_historical",
        },
        "optionalPartialPeriod": {
            "startDate": OPTIONAL_PARTIAL_START.isoformat(),
            "endDate": OPTIONAL_PARTIAL_END.isoformat(),
            "imported": False,
            "sessionGroups": len(optional_groups),
            "sourceAttendanceFlags": sum(
                1 for _, row in optional_rows for index in (4, 6)
                if row[index] not in (None, 0, "")
            ),
            "uniquePax": len(optional_names),
            "reasonExcluded": "Valid source data, but only a partial Monday-Tuesday period.",
        },
        "members": members,
        "sites": sorted(sites_by_address.values(), key=lambda item: item["name"].casefold()),
        "aos": sorted(aos, key=lambda item: item["name"].casefold()),
        "schedules": sorted(schedules, key=lambda item: (item["weekday"], item["startTime"], item["aoSourceKey"])),
        "sessions": sessions,
        "rawSourceRows": sorted(raw_source_rows, key=lambda item: item["rowNumber"]),
        "prospectiveSlots": sorted(prospective_slots, key=lambda item: (item["date"], item["startTime"], item["aoId"])),
    }
    manifest["summary"] = {
        "historicalSessions": len(sessions),
        "sourcePostFlags": sum(session["sourcePostFlags"] for session in sessions),
        "sourceQSourceFlags": sum(session["sourceQSourceFlags"] for session in sessions),
        "rawSourceAttendanceFlags": sum(session["sourceAttendanceFlags"] for session in sessions),
        "runtimeAttendanceRelationships": sum(session["runtimeAttendanceRelationships"] for session in sessions),
        "uniquePax": len(members),
        "uniqueQs": len({q_id for session in sessions for q_id in session["qIds"]}),
        "qAssignmentFlags": sum(len(session["qNames"]) for session in sessions),
        "fngRecords": sum(len(session["fngs"]) for session in sessions),
        "sites": len(manifest["sites"]),
        "aos": len(aos),
        "recurringSchedules": len(schedules),
        "prospectiveSlots": len(prospective_slots),
    }
    expected_summary = {
        "historicalSessions": 36,
        "sourcePostFlags": 239,
        "sourceQSourceFlags": 45,
        "rawSourceAttendanceFlags": 284,
        "runtimeAttendanceRelationships": 265,
        "uniquePax": 68,
        "uniqueQs": 20,
        "qAssignmentFlags": 36,
        "fngRecords": 4,
        "sites": 12,
        "aos": 18,
        "recurringSchedules": 18,
        "prospectiveSlots": 36,
    }
    # qNames is a runtime union. The raw four QSource-Q flags remain in raw rows.
    if manifest["summary"] != expected_summary:
        raise ValueError(
            "Corrected manifest control totals failed:\n"
            + json.dumps({"actual": manifest["summary"], "expected": expected_summary}, indent=2)
        )
    return manifest


def render_dry_run(manifest: dict) -> str:
    summary = manifest["summary"]
    slot_counts = defaultdict(int)
    for slot in manifest["prospectiveSlots"]:
        ao_name = next(ao["name"] for ao in manifest["aos"] if ao["id"] == slot["aoId"])
        slot_counts[ao_name] += 1
    lines = [
        "# F3 SacTown Demo Corrected Dry Run",
        "",
        "No historical dates are transposed. No historical facts are fabricated.",
        "",
        "## Historical controls",
        "",
        f"- Date range: {manifest['historicalPolicy']['startDate']} through {manifest['historicalPolicy']['endDate']}",
        f"- Sessions: {summary['historicalSessions']}",
        f"- Raw Post flags: {summary['sourcePostFlags']}",
        f"- Raw QSource flags: {summary['sourceQSourceFlags']}",
        f"- Raw attendance flags: {summary['rawSourceAttendanceFlags']}",
        f"- Runtime attendee relationships: {summary['runtimeAttendanceRelationships']}",
        f"- Unique PAX: {summary['uniquePax']}",
        f"- Unique runtime Qs: {summary['uniqueQs']}",
        f"- Runtime Q relationships: {summary['qAssignmentFlags']}",
        "- Raw Q/QSource-Q flags: 37 (preserved in raw staging provenance)",
        f"- FNG records: {summary['fngRecords']}",
        f"- Sites: {summary['sites']}",
        f"- AOs and recurring schedules: {summary['aos']}",
        "",
        "## Optional partial period excluded",
        "",
        "September 14-15 is valid source data but is not imported because it is only a partial Monday-Tuesday period.",
        "",
        "## Generated prospective Q slots",
        "",
        f"- Window: {SLOT_START.isoformat()} through {SLOT_END.isoformat()}",
        f"- Open slots: {summary['prospectiveSlots']}",
        "- Every slot is generated directly from an active source schedule.",
        "- No future Q assignments are created.",
        "",
    ]
    for ao_name in sorted(slot_counts, key=str.casefold):
        schedule = next(s for s in manifest["schedules"] if next(a["id"] for a in manifest["aos"] if a["name"] == ao_name) == s["aoId"])
        lines.append(f"- {ao_name}: weekday {schedule['weekday']}, {schedule['startTime']}, {schedule['durationMinutes']} minutes — {slot_counts[ao_name]} slots")
    return "\n".join(lines) + "\n"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=Path("import/sactown/output"))
    args = parser.parse_args()
    manifest = build_manifest(args.input)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = args.output_dir / "sactown_demo_manifest.json"
    report_path = args.output_dir / "sactown_demo_dry_run.md"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(render_dry_run(manifest), encoding="utf-8")
    print(json.dumps({"manifest": str(manifest_path), "report": str(report_path), "summary": manifest["summary"]}, indent=2))


if __name__ == "__main__":
    main()
