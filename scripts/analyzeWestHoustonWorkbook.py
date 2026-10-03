#!/usr/bin/env python3
"""Read-only, deterministic transformer for West Houston attendance workbooks.

Requires Python 3 and openpyxl. The source workbook is never modified.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import openpyxl


PARSER_VERSION = "west-houston-workbook-v1"
DERIVED_SHEETS = {"Posting Table", "FNG", "Combined"}
AO_MAP = {
    "the branch": "The Branch",
    "the point": "The Point",
    "the hop": "The HOP",
    "corridor": "The Corridor",
    "the corridor": "The Corridor",
    "the valley": "The Valley",
    "the tower": "The Tower",
    "valhalla": "Valhalla",
    "the oasis": "The Oasis",
    "the iron gate": "The Iron Gate",
    "the knot": "The Knot",
    "dr": "DR",
    "convergence": "Convergence",
}
SPECIAL_AO_TERMS = re.compile(
    r"\b(dr|converg|launch|annivers|closed|black\s*ops|otb|special|csaup)\b",
    re.IGNORECASE,
)


def text(value: Any) -> str:
    return " ".join(str(value or "").replace("\u00a0", " ").split())


def mechanical_name_key(value: Any) -> str:
    return text(value).casefold()


def loose_name_key(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", text(value).casefold())
    return re.sub(r"[^a-z0-9]+", "", normalized)


def slug(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.casefold())
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value or "unknown"


def stable_key(prefix: str, *parts: str) -> str:
    material = "|".join([PARSER_VERSION, prefix, *parts])
    digest = hashlib.sha256(material.encode("utf-8")).hexdigest()[:20]
    return f"wh-{prefix}-{digest}"


def iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 1 <= value <= 100000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    raw = text(value)
    if not raw:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def truthy_marker(value: Any) -> bool:
    if isinstance(value, (int, float)):
        return value != 0
    return text(value).casefold() in {"1", "true", "yes", "x"}


def normalize_bd_type(value: Any) -> str | None:
    raw = text(value)
    key = raw.casefold()
    aliases = {
        "bc": "Bootcamp",
        "bootcamp": "Bootcamp",
        "sb": "SB",
        "ruck": "Ruck",
        "run": "Run",
        "run/ruck": "Run/Ruck",
        "30/30": "30/30",
        "wod": "WOD",
        "dev": "Dev",
    }
    return aliases.get(key, raw or None)


def normalize_ao(value: Any, fallback: str) -> str:
    raw = text(value) or fallback
    return AO_MAP.get(raw.casefold(), raw)


def representative(items: list[dict[str, Any]], limit: int = 8) -> list[dict[str, Any]]:
    return items[:limit]


def analyze(workbook_path: Path) -> dict[str, Any]:
    workbook_bytes = workbook_path.read_bytes()
    workbook_hash = hashlib.sha256(workbook_bytes).hexdigest()
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)
    source_sheets = [name for name in workbook.sheetnames if name not in DERIVED_SHEETS]

    raw_rows_examined = 0
    nonempty_rows = 0
    rejected: list[dict[str, Any]] = []
    input_records: list[dict[str, Any]] = []
    ao_inventory: Counter[tuple[str, str]] = Counter()
    sheet_summaries: list[dict[str, Any]] = []

    for sheet_name in source_sheets:
        sheet = workbook[sheet_name]
        headers = [text(cell.value).casefold() for cell in sheet[1][:15]]

        def column(*names: str) -> int | None:
            for name in names:
                if name.casefold() in headers:
                    return headers.index(name.casefold())
            return None

        date_index = column("Date")
        name_index = column("Name")
        bd_type_index = column("BD Type")
        bd_index = column("BD")
        dd_index = column("DD")
        dd_type_index = column("DD Type")
        location_index = column("Location Comment")
        pax_comment_index = column("Pax Comment", "PAX Comment")
        home_index = column("Home AO or Visitor")
        bd_comment_index = column("BD Comment")
        if date_index is None or name_index is None or bd_index is None or dd_index is None:
            raise ValueError(f"Required source columns are missing from {sheet_name}")
        sheet_nonempty = 0
        sheet_candidates = 0

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            raw_rows_examined += 1
            values = list(row[:15]) + [None] * max(0, 15 - len(row))
            if not any(value not in (None, "") for value in values):
                continue
            nonempty_rows += 1
            sheet_nonempty += 1

            source_date = iso_date(values[date_index])
            raw_name = text(values[name_index])
            provenance = {"sheet": sheet_name, "row": row_number}

            if not source_date or not raw_name:
                rejected.append({
                    **provenance,
                    "reason": "missing_or_invalid_date_or_name",
                    "rawDate": text(values[date_index]),
                    "rawName": raw_name,
                })
                continue
            if mechanical_name_key(raw_name) in {"dd count", "bd count", "name", "pax"}:
                rejected.append({**provenance, "reason": "summary_or_header_row", "rawName": raw_name})
                continue

            is_bd = truthy_marker(values[bd_index])
            is_dd = truthy_marker(values[dd_index])
            if not is_bd and not is_dd:
                rejected.append({
                    **provenance,
                    "reason": "no_attendance_marker",
                    "rawName": raw_name,
                    "date": source_date,
                })
                continue

            sheet_candidates += 1
            raw_location = text(values[location_index]) if location_index is not None else ""
            raw_location = raw_location or sheet_name
            proposed_ao = normalize_ao(raw_location, sheet_name)
            ao_inventory[(raw_location, proposed_ao)] += 1
            pax_comment = text(values[pax_comment_index]) if pax_comment_index is not None else ""
            home_or_visitor = text(values[home_index]) if home_index is not None else ""
            bd_comment = text(values[bd_comment_index]) if bd_comment_index is not None else ""

            base = {
                "sourceSheet": sheet_name,
                "sourceRow": row_number,
                "date": source_date,
                "rawName": raw_name,
                "nameKey": mechanical_name_key(raw_name),
                "rawLocation": raw_location,
                "proposedAo": proposed_ao,
                "paxComment": pax_comment or None,
                "homeAoOrVisitor": home_or_visitor or None,
                "bdComment": bd_comment or None,
            }
            if is_bd:
                input_records.append({
                    **base,
                    "stream": "bd",
                    "workoutType": (
                        normalize_bd_type(values[bd_type_index])
                        if bd_type_index is not None
                        else None
                    ) or "Unspecified BD",
                    "isQ": pax_comment.casefold() in {"q", "vq"},
                    "qEvidence": pax_comment if pax_comment.casefold() in {"q", "vq"} else None,
                    "isFng": pax_comment.casefold() == "fng",
                })
            if is_dd:
                dd_type = normalize_bd_type(values[dd_type_index]) if dd_type_index is not None else None
                dd_type = dd_type or "Unspecified DD"
                input_records.append({
                    **base,
                    "stream": "dd",
                    "workoutType": dd_type,
                    "isQ": False,
                    "qEvidence": None,
                    "isFng": False,
                })

        sheet_summaries.append({
            "sheet": sheet_name,
            "physicalDataRows": max(sheet.max_row - 1, 0),
            "nonemptyRows": sheet_nonempty,
            "acceptedSourceRows": sheet_candidates,
            "dateColumn": date_index + 1,
            "nameColumn": name_index + 1,
        })

    variants_by_name: dict[str, Counter[str]] = defaultdict(Counter)
    member_dates: dict[str, list[str]] = defaultdict(list)
    member_provenance: dict[str, list[dict[str, Any]]] = defaultdict(list)
    explicit_fng_dates: dict[str, list[str]] = defaultdict(list)
    for record in input_records:
        key = record["nameKey"]
        variants_by_name[key][record["rawName"]] += 1
        member_dates[key].append(record["date"])
        member_provenance[key].append({"sheet": record["sourceSheet"], "row": record["sourceRow"]})
        if record["isFng"]:
            explicit_fng_dates[key].append(record["date"])

    members: list[dict[str, Any]] = []
    member_key_by_name: dict[str, str] = {}
    variant_collapses: list[dict[str, Any]] = []
    for name_key in sorted(variants_by_name):
        variants = variants_by_name[name_key]
        canonical_name = sorted(variants, key=lambda item: (-variants[item], item.casefold(), item))[0]
        dates = sorted(member_dates[name_key])
        fng_dates = sorted(set(explicit_fng_dates[name_key]))
        member_key = stable_key("member", name_key)
        member_key_by_name[name_key] = member_key
        member = {
            "memberKey": member_key,
            "proposedF3Name": canonical_name,
            "rawNameVariants": [
                {"value": variant, "occurrences": variants[variant]}
                for variant in sorted(variants, key=lambda item: (item.casefold(), item))
            ],
            "firstKnownAttendanceDate": dates[0],
            "lastKnownAttendanceDate": dates[-1],
            "explicitFngDates": fng_dates,
            "proposedFirstPostDate": fng_dates[0] if fng_dates else None,
            "firstPostSource": "explicit_fng" if fng_dates else None,
            "provenance": sorted(
                {f"{item['sheet']}!{item['row']}" for item in member_provenance[name_key]}
            ),
        }
        members.append(member)
        if len(variants) > 1:
            variant_collapses.append({
                "memberKey": member_key,
                "proposedF3Name": canonical_name,
                "variants": member["rawNameVariants"],
            })

    loose_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for member in members:
        loose_groups[loose_name_key(member["proposedF3Name"])].append(member)
    identity_collisions = [
        {
            "looseKey": key,
            "members": [
                {"memberKey": member["memberKey"], "proposedF3Name": member["proposedF3Name"]}
                for member in sorted(group, key=lambda item: item["proposedF3Name"].casefold())
            ],
            "reason": "punctuation_or_spacing_insensitive_names_collide",
        }
        for key, group in sorted(loose_groups.items())
        if key and len(group) > 1
    ]

    session_groups: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in input_records:
        group_key = (
            record["date"],
            record["proposedAo"],
            record["stream"],
            record["workoutType"],
        )
        session_groups[group_key].append(record)

    sessions: list[dict[str, Any]] = []
    attendance: list[dict[str, Any]] = []
    q_assignments: list[dict[str, Any]] = []
    fng_events: list[dict[str, Any]] = []
    duplicate_attendance: list[dict[str, Any]] = []
    contradictions: list[dict[str, Any]] = []

    for group_key in sorted(session_groups):
        session_date, proposed_ao, stream, workout_type = group_key
        records = sorted(session_groups[group_key], key=lambda item: (item["sourceSheet"], item["sourceRow"]))
        session_key = stable_key("session", session_date, proposed_ao.casefold(), stream, workout_type.casefold())
        raw_locations = sorted({record["rawLocation"] for record in records})
        comments = sorted({record["bdComment"] for record in records if record["bdComment"]})
        seen_members: dict[str, dict[str, Any]] = {}
        q_members: set[str] = set()
        fng_members: set[str] = set()
        for record in records:
            member_key = member_key_by_name[record["nameKey"]]
            if member_key in seen_members:
                duplicate_attendance.append({
                    "sessionKey": session_key,
                    "memberKey": member_key,
                    "kept": seen_members[member_key],
                    "removed": {"sheet": record["sourceSheet"], "row": record["sourceRow"]},
                })
            source = {"sheet": record["sourceSheet"], "row": record["sourceRow"]}
            if member_key not in seen_members:
                seen_members[member_key] = source
                attendance.append({"sessionKey": session_key, "memberKey": member_key, "provenance": source})
            if record["isQ"] and member_key not in q_members:
                q_members.add(member_key)
                q_assignments.append({
                    "sessionKey": session_key,
                    "memberKey": member_key,
                    "evidence": record["qEvidence"],
                    "provenance": source,
                })
            if record["isFng"] and member_key not in fng_members:
                fng_members.add(member_key)
                fng_events.append({"sessionKey": session_key, "memberKey": member_key, "provenance": source})
        for member_key in sorted(q_members & fng_members):
            contradictions.append({
                "type": "fng_and_q_same_session",
                "sessionKey": session_key,
                "memberKey": member_key,
                "provenance": [
                    {"sheet": record["sourceSheet"], "row": record["sourceRow"]}
                    for record in records
                    if member_key_by_name[record["nameKey"]] == member_key
                    and (record["isQ"] or record["isFng"])
                ],
            })
        sessions.append({
            "sessionKey": session_key,
            "date": session_date,
            "sourceLocations": raw_locations,
            "proposedAo": proposed_ao,
            "stream": stream,
            "bdType": workout_type,
            "bdComments": comments,
            "sourceSheets": sorted({record["sourceSheet"] for record in records}),
            "provenance": [{"sheet": record["sourceSheet"], "row": record["sourceRow"]} for record in records],
        })

    q_by_session = Counter(item["sessionKey"] for item in q_assignments)
    attendance_by_session = Counter(item["sessionKey"] for item in attendance)
    for session in sessions:
        session["attendanceCount"] = attendance_by_session[session["sessionKey"]]
        session["qCount"] = q_by_session[session["sessionKey"]]

    fng_discrepancies = [
        {
            "memberKey": member["memberKey"],
            "proposedF3Name": member["proposedF3Name"],
            "earliestAttendance": member["firstKnownAttendanceDate"],
            "earliestExplicitFng": member["explicitFngDates"][0],
        }
        for member in members
        if member["explicitFngDates"]
        and member["firstKnownAttendanceDate"] != member["explicitFngDates"][0]
    ]

    date_ao_sessions: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for session in sessions:
        date_ao_sessions[(session["date"], session["proposedAo"])].append(session)
    ambiguous_groupings = [
        {
            "date": key[0],
            "proposedAo": key[1],
            "reason": "multiple_workout_streams_or_types_same_date_and_ao",
            "sessions": [
                {"sessionKey": item["sessionKey"], "stream": item["stream"], "bdType": item["bdType"]}
                for item in group
            ],
        }
        for key, group in sorted(date_ao_sessions.items())
        if len(group) > 1
    ]

    unusual_aos = [
        {"rawLocation": raw, "proposedAo": proposed, "attendanceRecords": count}
        for (raw, proposed), count in sorted(ao_inventory.items())
        if proposed not in set(AO_MAP.values()) - {"DR", "Convergence"}
        or proposed in {"DR", "Convergence"}
    ]
    special_comments = []
    for session in sessions:
        matched = [comment for comment in session["bdComments"] if SPECIAL_AO_TERMS.search(comment)]
        if matched:
            special_comments.append({
                "sessionKey": session["sessionKey"],
                "date": session["date"],
                "proposedAo": session["proposedAo"],
                "comments": matched,
            })

    inviter_mentions = []
    seen_inviter_mentions: set[tuple[str, int]] = set()
    for record in input_records:
        marker = (record["sourceSheet"], record["sourceRow"])
        if marker in seen_inviter_mentions:
            continue
        if text(record["paxComment"]).casefold() in {"pp", "bb"} or re.search(
            r"\b(proud\s*papa|papa|pp|bb)\b", record["bdComment"] or "", re.IGNORECASE
        ):
            seen_inviter_mentions.add(marker)
            inviter_mentions.append({
                "sheet": record["sourceSheet"],
                "row": record["sourceRow"],
                "date": record["date"],
                "rawName": record["rawName"],
                "paxComment": record["paxComment"],
                "bdComment": record["bdComment"],
                "finding": "inviter-like notation without a reliable invited-member-to-inviter mapping",
            })

    sessions_by_year_ao = Counter((s["date"][:4], s["proposedAo"], s["stream"]) for s in sessions)
    session_lookup = {session["sessionKey"]: session for session in sessions}
    attendance_by_year_ao = Counter(
        (session_lookup[item["sessionKey"]]["date"][:4], session_lookup[item["sessionKey"]]["proposedAo"], session_lookup[item["sessionKey"]]["stream"])
        for item in attendance
    )

    output = {
        "metadata": {
            "parserVersion": PARSER_VERSION,
            "workbookFilename": workbook_path.name,
            "workbookSha256": workbook_hash,
            "sourceSheets": source_sheets,
            "excludedDerivedSheets": sorted(DERIVED_SHEETS),
            "sessionIdentityRule": "date + proposed AO + stream (BD/DD) + normalized workout type",
        },
        "summary": {
            "physicalRowsExamined": raw_rows_examined,
            "nonemptyRowsExamined": nonempty_rows,
            "legitimateSourceAttendanceRows": len({(r["sourceSheet"], r["sourceRow"]) for r in input_records}),
            "normalizedAttendanceRecords": len(attendance),
            "proposedMembers": len(members),
            "mechanicalVariantGroups": len(variant_collapses),
            "possibleIdentityCollisions": len(identity_collisions),
            "proposedSessions": len(sessions),
            "explicitQAssignments": len(q_assignments),
            "sessionsWithZeroQs": sum(1 for s in sessions if s["qCount"] == 0),
            "sessionsWithOneQ": sum(1 for s in sessions if s["qCount"] == 1),
            "sessionsWithMultipleQs": sum(1 for s in sessions if s["qCount"] > 1),
            "explicitFngEvents": len(fng_events),
            "membersWithExplicitFirstPost": sum(1 for m in members if m["proposedFirstPostDate"]),
            "fngEarliestAttendanceDiscrepancies": len(fng_discrepancies),
            "duplicateAttendanceRemoved": len(duplicate_attendance),
            "ambiguousSessionGroupings": len(ambiguous_groupings),
            "rejectedRows": len(rejected),
            "sourceContradictions": len(contradictions),
        },
        "sheetAnalysis": sheet_summaries,
        "aoInventory": [
            {"rawLocation": raw, "proposedAo": proposed, "sourceRecords": count}
            for (raw, proposed), count in sorted(ao_inventory.items())
        ],
        "members": members,
        "sessions": sessions,
        "attendance": attendance,
        "qAssignments": q_assignments,
        "fngEvidence": fng_events,
        "inviterRelationships": [],
        "review": {
            "mechanicalVariantGroups": variant_collapses,
            "possibleIdentityCollisions": identity_collisions,
            "fngEarliestAttendanceDiscrepancies": fng_discrepancies,
            "duplicateAttendanceRemoved": duplicate_attendance,
            "unusualAoValues": unusual_aos,
            "specialEventComments": special_comments,
            "inviterMentions": inviter_mentions,
            "ambiguousSessionGroupings": ambiguous_groupings,
            "rejectedRows": rejected,
            "sourceContradictions": contradictions,
            "inviterFinding": "No explicit inviter target field exists. PP/BB tags do not identify the inviter and were not converted into relationships.",
        },
        "aggregates": {
            "sessionsByYearAo": [
                {"year": key[0], "ao": key[1], "stream": key[2], "count": count}
                for key, count in sorted(sessions_by_year_ao.items())
            ],
            "attendanceByYearAo": [
                {"year": key[0], "ao": key[1], "stream": key[2], "count": count}
                for key, count in sorted(attendance_by_year_ao.items())
            ],
        },
    }
    return output


class UnionFind:
    def __init__(self, keys: list[str]) -> None:
        self.parent = {key: key for key in keys}

    def find(self, key: str) -> str:
        parent = self.parent[key]
        if parent != key:
            self.parent[key] = self.find(parent)
        return self.parent[key]

    def union(self, keys: list[str]) -> None:
        roots = sorted({self.find(key) for key in keys})
        if not roots:
            return
        root = roots[0]
        for other in roots[1:]:
            self.parent[other] = root


def apply_human_resolutions(normalized: dict[str, Any], manifest: dict[str, Any]) -> dict[str, Any]:
    """Apply explicit human decisions without changing the normalized source layer."""
    source = manifest["source"]
    metadata = normalized["metadata"]
    if source["attendanceWorkbookSha256"] != metadata["workbookSha256"]:
        raise ValueError("Human-resolution manifest does not match the attendance workbook SHA-256")
    if source["parserVersion"] != metadata["parserVersion"]:
        raise ValueError("Human-resolution manifest does not match the parser version")

    member_by_key = {item["memberKey"]: item for item in normalized["members"]}
    all_member_keys = sorted(member_by_key)
    union_find = UnionFind(all_member_keys)
    review_ids_by_member: dict[str, set[str]] = defaultdict(set)
    approved_names_by_member: dict[str, set[str]] = defaultdict(set)
    resolution_by_review_id: dict[str, str] = {}

    for decision in manifest["identityDecisions"]:
        keys = sorted(decision["sourceIdentityKeys"])
        if any(key not in member_by_key for key in keys):
            raise ValueError(f"Unknown source identity in {decision['reviewId']}")
        resolution_by_review_id[decision["reviewId"]] = decision["resolution"]
        for key in keys:
            review_ids_by_member[key].add(decision["reviewId"])
        if decision["resolution"] == "MERGE":
            union_find.union(keys)
            approved_name = text(decision.get("approvedCanonicalF3Name"))
            if not approved_name:
                raise ValueError(f"Missing approved canonical name for {decision['reviewId']}")
            for key in keys:
                approved_names_by_member[key].add(approved_name)
        elif decision["resolution"] not in {"KEEP_SEPARATE", "UNRESOLVED"}:
            raise ValueError(f"Unsupported resolution in {decision['reviewId']}")

    for decision in manifest.get("crossGroupIdentityDecisions", []):
        keys = sorted(decision["sourceIdentityKeys"])
        if decision["resolution"] != "MERGE":
            raise ValueError(f"Unsupported cross-group resolution in {decision['decisionId']}")
        union_find.union(keys)
        approved_name = text(decision.get("approvedCanonicalF3Name"))
        for key in keys:
            review_ids_by_member[key].add(decision["decisionId"])
            approved_names_by_member[key].add(approved_name)

    clusters: dict[str, list[str]] = defaultdict(list)
    for key in all_member_keys:
        clusters[union_find.find(key)].append(key)

    fng_decision_by_source_member = {
        item["sourceMemberKey"]: item for item in manifest["fngDiscrepancyDecisions"]
    }
    source_to_canonical_member: dict[str, str] = {}
    canonical_members: list[dict[str, Any]] = []
    for source_keys in sorted((sorted(keys) for keys in clusters.values()), key=lambda keys: keys):
        canonical_key = stable_key("canonical-member", *source_keys)
        for key in source_keys:
            source_to_canonical_member[key] = canonical_key
        source_members = [member_by_key[key] for key in source_keys]
        approved_names = sorted({name for key in source_keys for name in approved_names_by_member[key]}, key=lambda value: (value.casefold(), value))
        if len({name.casefold() for name in approved_names}) > 1:
            raise ValueError(f"Conflicting approved canonical names for {source_keys}: {approved_names}")
        canonical_name = approved_names[0] if approved_names else sorted(
            (member["proposedF3Name"] for member in source_members),
            key=lambda value: (value.casefold(), value),
        )[0]
        variants: Counter[str] = Counter()
        for member in source_members:
            for variant in member["rawNameVariants"]:
                variants[variant["value"]] += variant["occurrences"]
        earliest_attendance = min(member["firstKnownAttendanceDate"] for member in source_members)
        latest_attendance = max(member["lastKnownAttendanceDate"] for member in source_members)
        explicit_fng_dates = sorted({value for member in source_members for value in member["explicitFngDates"]})
        fng_reviews = [fng_decision_by_source_member[key] for key in source_keys if key in fng_decision_by_source_member]
        first_post_source = "earliest_verified_attendance"
        if any(item["decision"] == "Unsure" for item in fng_reviews):
            first_post_source = "earliest_verified_attendance_with_unresolved_later_fng"
        review_ids = sorted({value for key in source_keys for value in review_ids_by_member[key]})
        canonical_members.append({
            "canonicalMemberKey": canonical_key,
            "canonicalF3Name": canonical_name,
            "sourceMemberKeys": source_keys,
            "rawNameVariants": [
                {"value": value, "occurrences": variants[value]}
                for value in sorted(variants, key=lambda item: (item.casefold(), item))
            ],
            "firstKnownAttendanceDate": earliest_attendance,
            "lastKnownAttendanceDate": latest_attendance,
            "canonicalFirstPostDate": earliest_attendance,
            "firstPostSource": first_post_source,
            "explicitFngDates": explicit_fng_dates,
            "humanReviewIds": review_ids,
            "provenance": sorted({value for member in source_members for value in member["provenance"]}),
            "matchExistingHint": (
                {"strategy": "match_existing", "normalizedF3Name": "doge"}
                if canonical_name.casefold() == "doge"
                else None
            ),
        })

    workout_by_source_session: dict[str, dict[str, Any]] = {}
    for decision in manifest["workoutDecisions"]:
        for session_key in decision["sourceSessionKeys"]:
            if session_key in workout_by_source_session:
                raise ValueError(f"Source session appears in multiple workout decisions: {session_key}")
            workout_by_source_session[session_key] = decision

    source_session_by_key = {item["sessionKey"]: item for item in normalized["sessions"]}
    canonical_session_groups: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    excluded_source_sessions: list[dict[str, Any]] = []
    for session in normalized["sessions"]:
        if session["stream"] == "dd":
            excluded_source_sessions.append({
                "sessionKey": session["sessionKey"],
                "reason": "DD_EXCLUDED_BY_POLICY",
                "provenance": session["provenance"],
            })
            continue
        if session["proposedAo"] == "DR":
            excluded_source_sessions.append({
                "sessionKey": session["sessionKey"],
                "reason": "DR_UNRESOLVED_PENDING_AOQ",
                "provenance": session["provenance"],
            })
            continue
        decision = workout_by_source_session.get(session["sessionKey"])
        if decision and decision["canonicalRepresentation"]["action"] == "merge_source_sessions":
            group_key = ("human-workout", decision["reviewId"])
        else:
            group_key = ("source-session", session["sessionKey"])
        canonical_session_groups[group_key].append(session)

    source_to_canonical_session: dict[str, str] = {}
    canonical_sessions: list[dict[str, Any]] = []
    for group_key, source_sessions in sorted(canonical_session_groups.items()):
        source_sessions = sorted(source_sessions, key=lambda item: item["sessionKey"])
        decisions = {
            workout_by_source_session[item["sessionKey"]]["reviewId"]: workout_by_source_session[item["sessionKey"]]
            for item in source_sessions
            if item["sessionKey"] in workout_by_source_session
        }
        decision = next(iter(decisions.values())) if decisions else None
        if decision:
            representation = decision["canonicalRepresentation"]
            canonical_ao = representation["canonicalAo"]
            workout_type = representation["canonicalWorkoutType"]
            event_name = representation.get("canonicalEventName")
        else:
            canonical_ao = source_sessions[0]["proposedAo"]
            workout_type = source_sessions[0]["bdType"]
            event_name = None
        session_dates = {item["date"] for item in source_sessions}
        if len(session_dates) != 1:
            raise ValueError(f"Cannot combine source sessions from different dates: {group_key}")
        session_date = next(iter(session_dates))
        canonical_session_key = stable_key(
            "canonical-session",
            session_date,
            canonical_ao.casefold(),
            workout_type.casefold(),
            (event_name or "").casefold(),
        )
        for source_session in source_sessions:
            source_to_canonical_session[source_session["sessionKey"]] = canonical_session_key
        canonical_sessions.append({
            "canonicalSessionKey": canonical_session_key,
            "date": session_date,
            "canonicalAo": canonical_ao,
            "workoutType": workout_type,
            "eventName": event_name,
            "sourceSessionKeys": [item["sessionKey"] for item in source_sessions],
            "sourceLocations": sorted({value for item in source_sessions for value in item["sourceLocations"]}),
            "sourceSheets": sorted({value for item in source_sessions for value in item["sourceSheets"]}),
            "sourceComments": sorted({value for item in source_sessions for value in item["bdComments"]}),
            "humanReviewIds": sorted(decisions),
            "provenance": sorted(
                (value for item in source_sessions for value in item["provenance"]),
                key=lambda value: (value["sheet"], value["row"]),
            ),
        })

    def consolidate_records(
        records: list[dict[str, Any]],
        include_evidence: bool = False,
    ) -> list[dict[str, Any]]:
        grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        for record in records:
            if record["sessionKey"] not in source_to_canonical_session:
                continue
            key = (
                source_to_canonical_session[record["sessionKey"]],
                source_to_canonical_member[record["memberKey"]],
            )
            grouped[key].append(record)
        consolidated = []
        for (canonical_session_key, canonical_member_key), group in sorted(grouped.items()):
            source_member_keys = sorted({item["memberKey"] for item in group})
            item = {
                "canonicalSessionKey": canonical_session_key,
                "canonicalMemberKey": canonical_member_key,
                "sourceSessionKeys": sorted({value["sessionKey"] for value in group}),
                "sourceMemberKeys": source_member_keys,
                "humanReviewIds": sorted({
                    review_id
                    for source_member_key in source_member_keys
                    for review_id in review_ids_by_member[source_member_key]
                }),
                "provenance": sorted(
                    (value["provenance"] for value in group),
                    key=lambda value: (value["sheet"], value["row"]),
                ),
            }
            if include_evidence:
                item["evidence"] = sorted({value.get("evidence") for value in group if value.get("evidence")})
            consolidated.append(item)
        return consolidated

    canonical_attendance = consolidate_records(normalized["attendance"])
    canonical_q_assignments = consolidate_records(normalized["qAssignments"], include_evidence=True)
    canonical_fng_evidence = consolidate_records(normalized["fngEvidence"])

    attendance_by_session = Counter(item["canonicalSessionKey"] for item in canonical_attendance)
    q_by_session = Counter(item["canonicalSessionKey"] for item in canonical_q_assignments)
    fng_by_session = Counter(item["canonicalSessionKey"] for item in canonical_fng_evidence)
    for session in canonical_sessions:
        key = session["canonicalSessionKey"]
        session["attendanceCount"] = attendance_by_session[key]
        session["qCount"] = q_by_session[key]
        session["fngCount"] = fng_by_session[key]

    q_pairs = {(item["canonicalSessionKey"], item["canonicalMemberKey"]) for item in canonical_q_assignments}
    fng_pairs = {(item["canonicalSessionKey"], item["canonicalMemberKey"]) for item in canonical_fng_evidence}
    canonical_contradictions = [
        {"canonicalSessionKey": session_key, "canonicalMemberKey": member_key, "type": "Q_AND_FNG_SAME_SESSION"}
        for session_key, member_key in sorted(q_pairs & fng_pairs)
    ]

    canonical_member_by_key = {item["canonicalMemberKey"]: item for item in canonical_members}
    fng_discrepancy_representations = []
    for decision in manifest["fngDiscrepancyDecisions"]:
        canonical_member_key = source_to_canonical_member[decision["sourceMemberKey"]]
        canonical_member = canonical_member_by_key[canonical_member_key]
        fng_discrepancy_representations.append({
            "reviewId": decision["reviewId"],
            "decision": decision["decision"],
            "canonicalMemberKey": canonical_member_key,
            "canonicalF3Name": canonical_member["canonicalF3Name"],
            "earliestRecordedAttendance": decision["earliestAttendance"],
            "laterExplicitFng": decision["explicitFng"],
            "canonicalFirstPostDate": canonical_member["canonicalFirstPostDate"],
            "representation": "Earlier verified attendance remains canonical first-post history; later explicit FNG evidence is retained without overwriting it.",
        })

    unresolved_identity_decisions = [
        item for item in manifest["identityDecisions"] if item["resolution"] == "UNRESOLVED"
    ]
    doge_members = [item for item in canonical_members if item["matchExistingHint"]]
    if len(doge_members) != 1:
        raise ValueError(f"Expected exactly one canonical DOGE member, found {len(doge_members)}")

    session_keys = [item["canonicalSessionKey"] for item in canonical_sessions]
    member_keys = [item["canonicalMemberKey"] for item in canonical_members]
    if len(session_keys) != len(set(session_keys)):
        raise ValueError("Duplicate canonical session keys")
    if len(member_keys) != len(set(member_keys)):
        raise ValueError("Duplicate canonical member keys")

    eligible_source_session_keys = {
        item["sessionKey"]
        for item in normalized["sessions"]
        if item["stream"] == "bd" and item["proposedAo"] != "DR"
    }
    excluded_dd_session_keys = {
        item["sessionKey"] for item in excluded_source_sessions if item["reason"] == "DD_EXCLUDED_BY_POLICY"
    }
    excluded_dr_session_keys = {
        item["sessionKey"] for item in excluded_source_sessions if item["reason"] == "DR_UNRESOLVED_PENDING_AOQ"
    }

    def count_records_for_sessions(records: list[dict[str, Any]], session_keys: set[str]) -> int:
        return sum(1 for item in records if item["sessionKey"] in session_keys)

    before_canonical_scope = {
        "members": len(normalized["members"]),
        "primarySessionsExcludingDr": len(eligible_source_session_keys),
        "primaryAttendanceExcludingDr": count_records_for_sessions(normalized["attendance"], eligible_source_session_keys),
        "primaryQAssignmentsExcludingDr": count_records_for_sessions(normalized["qAssignments"], eligible_source_session_keys),
        "primaryFngEvidenceExcludingDr": count_records_for_sessions(normalized["fngEvidence"], eligible_source_session_keys),
    }

    return {
        "metadata": {
            "schemaVersion": "west-houston-post-human-canonical-v1",
            "parserVersion": metadata["parserVersion"],
            "attendanceWorkbookFilename": metadata["workbookFilename"],
            "attendanceWorkbookSha256": metadata["workbookSha256"],
            "humanReviewWorkbookFilename": source["humanReviewWorkbookFilename"],
            "humanReviewWorkbookSha256": source["humanReviewWorkbookSha256"],
            "humanResolutionManifestSchemaVersion": manifest["schemaVersion"],
            "canonicalSessionPolicy": "Primary BD only; DD excluded; DR excluded pending AOQ; explicit workout decisions applied.",
        },
        "beforeHumanReview": normalized["summary"],
        "beforeCanonicalScope": before_canonical_scope,
        "summary": {
            "canonicalMembers": len(canonical_members),
            "approvedMergeGroups": sum(1 for item in manifest["identityDecisions"] if item["resolution"] == "MERGE"),
            "approvedKeepSeparateGroups": sum(1 for item in manifest["identityDecisions"] if item["resolution"] == "KEEP_SEPARATE"),
            "unresolvedIdentityGroups": len(unresolved_identity_decisions),
            "canonicalPrimarySessions": len(canonical_sessions),
            "canonicalAttendance": len(canonical_attendance),
            "canonicalQAssignments": len(canonical_q_assignments),
            "canonicalFngEvidence": len(canonical_fng_evidence),
            "unresolvedFngDiscrepancies": sum(1 for item in manifest["fngDiscrepancyDecisions"] if item["decision"] == "Unsure"),
            "canonicalContradictions": len(canonical_contradictions),
            "excludedDdSessions": sum(1 for item in excluded_source_sessions if item["reason"] == "DD_EXCLUDED_BY_POLICY"),
            "excludedDrSessions": sum(1 for item in excluded_source_sessions if item["reason"] == "DR_UNRESOLVED_PENDING_AOQ"),
            "excludedDdAttendance": count_records_for_sessions(normalized["attendance"], excluded_dd_session_keys),
            "excludedDrAttendance": count_records_for_sessions(normalized["attendance"], excluded_dr_session_keys),
            "excludedDdQAssignments": count_records_for_sessions(normalized["qAssignments"], excluded_dd_session_keys),
            "excludedDrQAssignments": count_records_for_sessions(normalized["qAssignments"], excluded_dr_session_keys),
            "excludedDdFngEvidence": count_records_for_sessions(normalized["fngEvidence"], excluded_dd_session_keys),
            "excludedDrFngEvidence": count_records_for_sessions(normalized["fngEvidence"], excluded_dr_session_keys),
        },
        "appliedHumanReview": {
            "reviewer": source["reviewer"],
            "identityDecisions": manifest["identityDecisions"],
            "crossGroupIdentityDecisions": manifest.get("crossGroupIdentityDecisions", []),
            "fngDiscrepancyDecisions": manifest["fngDiscrepancyDecisions"],
            "workoutDecisions": manifest["workoutDecisions"],
            "streamPolicies": manifest["streamPolicies"],
        },
        "canonicalMembers": canonical_members,
        "canonicalSessions": canonical_sessions,
        "canonicalAttendance": canonical_attendance,
        "canonicalQAssignments": canonical_q_assignments,
        "canonicalFngEvidence": canonical_fng_evidence,
        "fngDiscrepancyRepresentations": fng_discrepancy_representations,
        "excludedSourceSessions": excluded_source_sessions,
        "remainingUnresolvedIdentityGroups": unresolved_identity_decisions,
        "canonicalContradictions": canonical_contradictions,
        "existingMemberMatchCandidates": doge_members,
        "normalizedSource": normalized,
    }


def post_human_markdown_report(result: dict[str, Any]) -> str:
    summary = result["summary"]
    before = result["beforeHumanReview"]
    before_scope = result["beforeCanonicalScope"]
    identity_decisions = result["appliedHumanReview"]["identityDecisions"]
    unresolved_identity = result["remainingUnresolvedIdentityGroups"]
    unresolved_fng = result["fngDiscrepancyRepresentations"]
    lines = [
        "# West Houston post-human-review canonical dry run",
        "",
        "## Result",
        "",
        "DOGE's explicit decisions were applied to the normalized workbook evidence. No import or database write was performed.",
        "",
        "## Before and after",
        "",
        "| Measure | Before human review | Post-review canonical | Change |",
        "| --- | ---: | ---: | ---: |",
        f"| Members | {before['proposedMembers']} | {summary['canonicalMembers']} | {summary['canonicalMembers'] - before['proposedMembers']} |",
        f"| Identity collisions requiring a decision | {before['possibleIdentityCollisions']} | {summary['unresolvedIdentityGroups']} | {summary['unresolvedIdentityGroups'] - before['possibleIdentityCollisions']} |",
        f"| Primary sessions eligible before human decisions | {before_scope['primarySessionsExcludingDr']} | {summary['canonicalPrimarySessions']} | {summary['canonicalPrimarySessions'] - before_scope['primarySessionsExcludingDr']} |",
        f"| Primary attendance eligible before human decisions | {before_scope['primaryAttendanceExcludingDr']} | {summary['canonicalAttendance']} | {summary['canonicalAttendance'] - before_scope['primaryAttendanceExcludingDr']} |",
        f"| Primary Q/VQ assignments eligible before human decisions | {before_scope['primaryQAssignmentsExcludingDr']} | {summary['canonicalQAssignments']} | {summary['canonicalQAssignments'] - before_scope['primaryQAssignmentsExcludingDr']} |",
        f"| Primary explicit FNG evidence eligible before human decisions | {before_scope['primaryFngEvidenceExcludingDr']} | {summary['canonicalFngEvidence']} | {summary['canonicalFngEvidence'] - before_scope['primaryFngEvidenceExcludingDr']} |",
        f"| Unresolved attendance-before-FNG cases | {before['fngEarliestAttendanceDiscrepancies']} | {summary['unresolvedFngDiscrepancies']} | 0 |",
        "",
        f"The complete normalized source layer still contains {before['proposedSessions']} BD/DD sessions and {before['normalizedAttendanceRecords']} attendance records. The comparable before column above excludes DD and unresolved DR before measuring the effects of human decisions.",
        "",
        "## Identity reconciliation",
        "",
        f"- MERGE: {sum(1 for item in identity_decisions if item['resolution'] == 'MERGE')} review groups.",
        f"- KEEP_SEPARATE: {sum(1 for item in identity_decisions if item['resolution'] == 'KEEP_SEPARATE')} review groups.",
        f"- UNRESOLVED: {sum(1 for item in identity_decisions if item['resolution'] == 'UNRESOLVED')} review group.",
        "- WH-ID-034 and WH-ID-035 were joined by the explicit cross-group note `Same as above` and their shared approved name `lightningrod`.",
        f"- Canonical member count: {summary['canonicalMembers']}.",
        "- DOGE remains a distinct canonical member with a `match_existing` hint for eventual adapter matching.",
        "",
        "## Workout reconciliation",
        "",
        "- The Knot on 2026-06-13 is one Bootcamp including Java.",
        "- The Knot on 2026-06-25 is one Bootcamp including Gladiator.",
        "- The 2025-07-12 rows whose source location is `Convergence` are represented as `The Corridor 4-year anniversary` at The Corridor. The original location and workbook rows remain in provenance.",
        "- No permanent Convergence AO is created.",
        "",
        "## FNG handling",
        "",
        f"All {len(unresolved_fng)} decisions remain `Unsure`. For every case, the earlier verified attendance date remains canonical first-post history. The later explicit FNG record remains separate evidence with its original AO and workbook-row provenance.",
        "",
        "| Review ID | F3 name | Earliest attendance | Later explicit FNG | Canonical first post |",
        "| --- | --- | --- | --- | --- |",
    ]
    for item in unresolved_fng:
        earliest = item["earliestRecordedAttendance"]
        later = item["laterExplicitFng"]
        lines.append(
            f"| {item['reviewId']} | {item['canonicalF3Name']} | {earliest['date']} at {earliest['ao']} | {later['date']} at {later['ao']} | {item['canonicalFirstPostDate']} |"
        )
    lines += [
        "",
        "## DD and DR exclusion checks",
        "",
        f"- Excluded DD source sessions: {summary['excludedDdSessions']}.",
        f"- Excluded DR source sessions: {summary['excludedDrSessions']}.",
        f"- Excluded DD attendance records: {summary['excludedDdAttendance']}; excluded DR attendance records: {summary['excludedDrAttendance']}.",
        f"- Excluded DD/DR Q assignments: {summary['excludedDdQAssignments'] + summary['excludedDrQAssignments']}; excluded DD/DR FNG records: {summary['excludedDdFngEvidence'] + summary['excludedDrFngEvidence']}.",
        "- Canonical sessions, attendance, Q assignments, FNG totals, and aggregate-ready records contain neither DD nor DR.",
        "- DD and DR source sessions remain listed with row provenance in `excludedSourceSessions`.",
        "",
        "## Validation",
        "",
        f"- Unique canonical member keys: {summary['canonicalMembers']}.",
        f"- Unique canonical session keys: {summary['canonicalPrimarySessions']}.",
        f"- Canonical Q/FNG contradictions: {summary['canonicalContradictions']}.",
        "- Every canonical attendance, Q, and FNG record retains source session keys, source member keys, workbook rows, and applicable human review IDs.",
        "",
        "## Preserved unresolved decisions (not import blockers)",
        "",
    ]
    for item in unresolved_identity:
        lines.append(
            f"- {item['reviewId']} ({' / '.join(item['candidateNames'])}): DOGE selected UNRESOLVED. Preserve the identities independently; no merge is approved."
        )
    lines.append(f"- WH-FNG-001 through WH-FNG-{len(unresolved_fng):03d}: DOGE selected Unsure. Earlier verified attendance remains first-post history and later explicit FNG evidence remains preserved; no further inference is approved.")
    lines.append("- WH-KNOW-001 (DR): AOQ confirmation is pending. DR remains outside this approved canonical import and does not block it.")
    lines.append("")
    return "\n".join(lines)


def markdown_report(result: dict[str, Any]) -> str:
    metadata = result["metadata"]
    summary = result["summary"]
    review = result["review"]

    def table(rows: list[dict[str, Any]], columns: list[tuple[str, str]]) -> list[str]:
        if not rows:
            return ["None."]
        lines = [
            "| " + " | ".join(label for _, label in columns) + " |",
            "| " + " | ".join("---" for _ in columns) + " |",
        ]
        for row in rows:
            values = []
            for key, _ in columns:
                value = row.get(key, "")
                if isinstance(value, (list, dict)):
                    value = json.dumps(value, ensure_ascii=False, sort_keys=True)
                values.append(str(value).replace("|", "\\|").replace("\n", " "))
            lines.append("| " + " | ".join(values) + " |")
        return lines

    lines = [
        "# West Houston historical workbook dry run",
        "",
        "## Source selection",
        "",
        f"- Workbook: `{metadata['workbookFilename']}`",
        f"- SHA-256: `{metadata['workbookSha256']}`",
        f"- Parser version: `{metadata['parserVersion']}`",
        f"- Canonical source sheets: {', '.join(metadata['sourceSheets'])}",
        f"- Excluded derived sheets: {', '.join(metadata['excludedDerivedSheets'])}",
        "- Reason: the AO sheets hold the imported row-level records. `Combined` unions those sheets, while `Posting Table` and `FNG` are calculated summaries.",
        "- Session identity: date + proposed AO + BD/DD stream + normalized workout type.",
        "",
        "## Reconciliation totals",
        "",
        *table([{"metric": key, "value": value} for key, value in summary.items()], [("metric", "Metric"), ("value", "Value")]),
        "",
        "## Source sheets",
        "",
        *table(result["sheetAnalysis"], [("sheet", "Sheet"), ("physicalDataRows", "Physical rows"), ("nonemptyRows", "Nonempty rows"), ("acceptedSourceRows", "Accepted source rows"), ("dateColumn", "Date column"), ("nameColumn", "Name column")]),
        "",
        "Valhalla is the column-order exception: its header and cached rows place Name in G and Date in H; the other AO sheets place Date in G and Name in H.",
        "",
        "## AO/location inventory",
        "",
        *table(result["aoInventory"], [("rawLocation", "Raw location"), ("proposedAo", "Proposed AO"), ("sourceRecords", "Accepted records")]),
        "",
        "## Sessions by year and AO",
        "",
        *table(result["aggregates"]["sessionsByYearAo"], [("year", "Year"), ("ao", "AO"), ("stream", "Stream"), ("count", "Sessions")]),
        "",
        "## Attendance by year and AO",
        "",
        *table(result["aggregates"]["attendanceByYearAo"], [("year", "Year"), ("ao", "AO"), ("stream", "Stream"), ("count", "Attendance")]),
        "",
        "## Review findings",
        "",
        "### Possible identity collisions",
        "",
        *table(representative(review["possibleIdentityCollisions"]), [("looseKey", "Loose key"), ("members", "Proposed members"), ("reason", "Reason")]),
        "",
        "### Mechanical name-variant collapses",
        "",
        *table(representative(review["mechanicalVariantGroups"]), [("proposedF3Name", "Proposed name"), ("variants", "Variants")]),
        "",
        "### FNG date discrepancies",
        "",
        *table(representative(review["fngEarliestAttendanceDiscrepancies"]), [("proposedF3Name", "Member"), ("earliestAttendance", "Earliest attendance"), ("earliestExplicitFng", "Earliest explicit FNG")]),
        "",
        "### Duplicate attendance removed",
        "",
        *table(representative(review["duplicateAttendanceRemoved"]), [("sessionKey", "Session"), ("memberKey", "Member"), ("kept", "Kept"), ("removed", "Removed")]),
        "",
        "### Unusual AO values",
        "",
        *table(representative(review["unusualAoValues"]), [("rawLocation", "Raw location"), ("proposedAo", "Proposed AO"), ("attendanceRecords", "Records")]),
        "",
        "### Special-event comments",
        "",
        *table(representative(review["specialEventComments"]), [("date", "Date"), ("proposedAo", "AO"), ("comments", "Comments")]),
        "",
        "### Ambiguous session groupings",
        "",
        *table(representative(review["ambiguousSessionGroupings"]), [("date", "Date"), ("proposedAo", "AO"), ("sessions", "Candidate sessions")]),
        "",
        "### Rejected rows",
        "",
        *table(representative(review["rejectedRows"]), [("sheet", "Sheet"), ("row", "Row"), ("reason", "Reason"), ("rawName", "Name"), ("rawDate", "Date")]),
        "",
        "### Source contradictions",
        "",
        *table(representative(review["sourceContradictions"]), [("type", "Type"), ("sessionKey", "Session"), ("memberKey", "Member"), ("provenance", "Source")]),
        "",
        "### Q-assignment examples",
        "",
        *table(representative(result["qAssignments"]), [("sessionKey", "Session"), ("memberKey", "Member"), ("evidence", "Evidence"), ("provenance", "Source")]),
        "",
        "### Multiple-Q session examples",
        "",
        *table(representative([item for item in result["sessions"] if item["qCount"] > 1]), [("date", "Date"), ("proposedAo", "AO"), ("stream", "Stream"), ("bdType", "Type"), ("qCount", "Q count")]),
        "",
        "### Explicit FNG examples",
        "",
        *table(representative(result["fngEvidence"]), [("sessionKey", "Session"), ("memberKey", "Member"), ("provenance", "Source")]),
        "",
        "### Inviter/Proud Papa",
        "",
        review["inviterFinding"],
        "",
        *table(representative(review["inviterMentions"]), [("sheet", "Sheet"), ("row", "Row"), ("date", "Date"), ("rawName", "Name"), ("paxComment", "PAX comment"), ("bdComment", "BD comment"), ("finding", "Finding")]),
        "",
        "## Decisions requiring review",
        "",
        "- Confirm that DD markers represent separate historical workout sessions rather than metadata on the primary BD. They are kept as separate `dd` sessions so same-date activity is not lost.",
        "- Review every punctuation-insensitive identity collision before merging people.",
        "- Review FNG records whose explicit FNG date is later than earliest attendance. The output retains the explicit FNG date and does not rewrite earlier attendance.",
        "- Confirm whether `DR` and `Convergence` should remain distinct AOs, map to event locations, or carry a separate event classification.",
        "- Review launch, anniversary, closure, OTB, Black Ops, and convergence comments before normalizing special events.",
        "- Decide whether `VQ` should remain Q evidence. This parser treats it as explicit Q evidence.",
        "- No inviter relationships are proposed because PP/BB markers do not identify a target inviter.",
        "",
        "## Current region-import framework differences",
        "",
        "- Staged session participants currently support only attendee, Q, and co-Q roles. This output stores FNG evidence separately while also keeping the member in attendance.",
        "- The current session commit RPC writes `sessions.fngs` as an empty array, so it cannot preserve this output's explicit FNG events.",
        "- The framework has no first-class fields for BD/DD stream, workout type, BD comment, raw location variants, or row-level provenance on committed sessions.",
        "- The framework can stage notes, but flattening structured comments and provenance into notes would lose semantics.",
        "- The current identity normalizer does not directly accept this output's complete name-variant and collision-review structure.",
        "- No write mapping has been implemented in this phase.",
    ]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--json-output", type=Path, required=True)
    parser.add_argument("--report-output", type=Path, required=True)
    parser.add_argument("--human-resolution-manifest", type=Path)
    args = parser.parse_args()

    normalized = analyze(args.workbook.resolve())
    if args.human_resolution_manifest:
        manifest = json.loads(args.human_resolution_manifest.read_text(encoding="utf-8"))
        result = apply_human_resolutions(normalized, manifest)
        report = post_human_markdown_report(result)
    else:
        result = normalized
        report = markdown_report(result)
    args.json_output.parent.mkdir(parents=True, exist_ok=True)
    args.report_output.parent.mkdir(parents=True, exist_ok=True)
    args.json_output.write_text(json.dumps(result, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
    args.report_output.write_text(report, encoding="utf-8")
    print(json.dumps(result["summary"], indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
